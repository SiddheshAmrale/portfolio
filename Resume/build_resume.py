"""Builds the one-page resume (DOCX + PDF) for Data, Platform and Forward Deployed roles.

    python build_resume.py

Outputs, next to this script:
    Siddhesh_Amrale_Resume.docx / .pdf  (the PDF is also copied to public/ for the site)
    Siddhesh_Amrale_Resume.txt          (for application forms that take pasted text)

Layout rules (ATS): one column, standard section names, no tables, text
boxes, icons or images; contact details in the body, not the page header.
The build fails if the PDF runs past one page.
"""

import html
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")

NAME = "Siddhesh Amrale"
HEADLINE = "Data & AI Platform Engineer | Pittsburgh, PA, open to relocation"
SITE = "https://gentle-river-08010bc0f.1.azurestaticapps.net"

# (shown text, link). Two lines, so no URL is split across a line break.
CONTACT = [
    [("siddhesh.amrale@gmail.com", "mailto:siddhesh.amrale@gmail.com"),
     ("+1 (973) 438-8729", "tel:+19734388729"),
     ("linkedin.com/in/siddheshnitinamrale", "https://www.linkedin.com/in/siddheshnitinamrale/")],
    [("github.com/SiddheshAmrale", "https://github.com/SiddheshAmrale"),
     ("Portfolio: gentle-river-08010bc0f.1.azurestaticapps.net", SITE)],
]

SUMMARY = (
    "Data and AI platform engineer, 6+ years. Started in full-stack product work in Pune, moved "
    "to mortgage data pipelines at Freddie Mac while earning an MS at NJIT, and now own Pittsburgh "
    "International Airport's data and AI platform. The loop: sit with the people who use the "
    "data, ship a working version fast, measure it, then harden it into a platform."
)

SKILLS = [
    ("Data", "SQL, Python, PySpark, Apache Spark, Databricks (Lakeflow/DLT, Unity Catalog), "
             "Delta Lake, dbt, Airflow, Snowflake, DuckDB, Parquet, ETL/ELT, CDC, SCD Type 2, "
             "data modeling, data warehousing, data quality, governance"),
    ("Streaming", "Kafka, Azure Event Hubs, IoT Hub, Azure Data Explorer (KQL), event-driven "
                  "architecture, late and out-of-order data"),
    ("AI", "LLM apps and pipelines, RAG, AI agents, MCP, LLM evaluation, Azure OpenAI, "
           "Semantic Kernel, prompt versioning, MLflow"),
    ("Cloud & platform", "Azure, AWS (S3, Lambda, EKS), Terraform, Bicep, Docker, Kubernetes, "
                         "GitHub Actions, Azure DevOps, ArgoCD, Prometheus, Grafana"),
    ("Languages & apps", "Python, SQL/KQL, C#/.NET, TypeScript, JavaScript/Node.js, Java, Rust, "
                         "React, Angular, REST, GraphQL"),
]

JOBS = [
    (
        "Analytics Full Stack Developer (Data & AI Platform Engineering)", "Jan 2024 – Present",
        "Pittsburgh International Airport (Allegheny County Airport Authority), Pittsburgh, PA",
        [
            "Designed and built the One Data Platform on Databricks: 60 Lakeflow (DLT) pipelines "
            "and 58 jobs on Unity Catalog across flights, baggage, parking, finance, workforce and "
            "customer experience; SCD2 silver and gold data products in Power BI for 5 business units.",
            "Built a reusable ingestion platform on .NET 8 Azure Functions (~50 connectors, 1,300+ "
            "unit tests) with watermark/CDC loads, gap detection and replay. New sources take 2 days "
            "instead of weeks; leaving Azure Data Factory identified $400K+ a year in savings.",
            "Partnered with operations teams to map how they really work, then shipped React + "
            "MSAL apps and Kusto-backed APIs they use daily (parking kiosks, live gate map, "
            "concessions); incident response time down 40%.",
            "Put LLMs inside the pipeline: Databricks ai_query (Llama 3.3 70B) turns FAA NOTAMs and "
            "passenger feedback into summaries and sentiment, with strict JSON output, versioned "
            "prompts, a kill switch and a content-hashed cache for deterministic reruns.",
            "Built RAG agents on Azure OpenAI and Semantic Kernel over governed data, gated by "
            "evaluation sets; five teams now self-serve answers and ad-hoc report requests fell ~60%.",
            "Built the real-time path: ~70 Azure Functions over IoT Hub and Event Hubs into Azure Data "
            "Explorer at ~5-minute freshness, all deployed with Terraform, Asset Bundles and Azure DevOps.",
        ],
    ),
    (
        "Software Engineer", "Sep 2022 – Nov 2023", "Freddie Mac, New York, NY (remote)",
        [
            "Built Python/SQL pipelines for loan-level mortgage data with automated validation and "
            "reconciliation; manual review down 40%.",
            "Designed a GraphQL + REST data gateway (Node.js, Java, Azure API Management, Entra ID) "
            "serving decisioning data at 1K+ requests/s, P95 under 10 ms; SLOs, circuit breakers "
            "and idempotent retries held 99.95% uptime.",
            "Moved 12+ services onto Kubernetes (AKS, EKS) with GitHub Actions and ArgoCD canary "
            "releases: release cycles 50% shorter, deployment incidents down 30%.",
        ],
    ),
    (
        "Software Engineer, Full-Stack", "Jan 2019 – Jul 2021",
        "Top Hat Software Technologies, Pune, India",
        [
            "Built full-stack apps (Angular, Java, MongoDB) for media delivery over Azure Blob "
            "Storage and CDN: 10M+ requests a day under 90 ms p99. Batch ETL cut metadata refresh "
            "from 4 hours to under 20 minutes.",
            "Added Prometheus and Grafana SLI dashboards and alerts, then fixed the worst "
            "bottlenecks: availability rose from 99.5% to 99.9%.",
        ],
    ),
]

PROJECTS = [
    ("Agent-Ready Lakehouse",
     "MCP server giving AI agents governed, audited, read-only access to lakehouse data; dbt "
     "layer with 16 data tests; CI eval gate on accuracy and 23 attack cases.",
     ("github.com/SiddheshAmrale/agent-ready-lakehouse",
      "https://github.com/SiddheshAmrale/agent-ready-lakehouse")),
    ("VELA",
     "Rust + libmpv media player: 5.6% CPU on HEVC vs 9.8% for stock mpv, 2.2 MB binary.",
     ("github.com/SiddheshAmrale/vela", "https://github.com/SiddheshAmrale/vela")),
    ("Lakehouse Forge",
     "Medallion lakehouse with CDC/SCD2, data contracts and freshness SLAs (Python, DuckDB).",
     ("portfolio /labs/forge", SITE + "/labs/forge")),
]

CERTS = ("Microsoft Certified: Azure AI Engineer Associate (2025) · Azure Developer "
         "Associate (2024) · Azure Administrator Associate (2022)")

EDUCATION = [
    ("M.S. Computer Science, New Jersey Institute of Technology", "2023"),
    ("B.E. Computer Engineering, Savitribai Phule Pune University", "2019"),
]


# ---------------------------------------------------------------- DOCX

INK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x44, 0x44, 0x44)
BODY_PT = 9.8


LINK = RGBColor(0x1F, 0x4E, 0x8C)


def _run(paragraph, text, size=None, color=None):
    r = paragraph.add_run(text)
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return r


def _hyperlink(paragraph, url, text, size=None):
    """A real, clickable hyperlink run (python-docx has no public API for this)."""
    rel = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), str(LINK))
    props.append(color)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        props.append(sz)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _border_bottom(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, val in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
        bottom.set(qn(k), val)
    bdr.append(bottom)
    p_pr.append(bdr)


def build_docx(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.55)
    sec.top_margin = sec.bottom_margin = Inches(0.45)
    right_tab = sec.page_width - sec.left_margin - sec.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(BODY_PT)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    def para(text="", size=None, bold=False, color=None, after=0, before=0, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        return p

    def heading(text):
        p = para(text.upper(), size=10.2, bold=True, before=5, after=2)
        _border_bottom(p)

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.add_run(text)

    def dated(left, right, before=4):
        p = para(before=before)
        p.paragraph_format.tab_stops.add_tab_stop(right_tab, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(left).bold = True
        p.add_run("\t" + right).bold = True

    para(NAME, size=19, bold=True)
    para(HEADLINE, size=11, color=MUTED)
    for line in CONTACT:
        p = para()
        for i, (text, url) in enumerate(line):
            if i:
                _run(p, " | ", size=9, color=MUTED)
            _hyperlink(p, url, text, size=9)

    heading("Summary")
    para(SUMMARY)

    heading("Skills")
    for group, items in SKILLS:
        p = para(after=0.5)
        p.add_run(group + ": ").bold = True
        p.add_run(items)

    heading("Experience")
    for title, dates, org, bullets in JOBS:
        dated(title, dates)
        para(org, italic=True, color=MUTED, after=1)
        for b in bullets:
            bullet(b)

    heading("Projects")
    for name, desc, where in PROJECTS:
        p = para(after=1)
        p.add_run(name + ": ").bold = True
        p.add_run(desc + " ")
        _hyperlink(p, where[1], where[0])

    heading("Education & certifications")
    for degree, year in EDUCATION:
        p = para()
        p.paragraph_format.tab_stops.add_tab_stop(right_tab, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(degree)
        p.add_run("\t" + year)

    para(CERTS, before=1)

    doc.core_properties.author = NAME
    doc.core_properties.title = f"{NAME} Resume"
    doc.save(path)


# ---------------------------------------------------------------- PDF (via HTML)

CSS = f"""
@page {{ size: Letter; margin: 0.45in 0.55in; }}
* {{ box-sizing: border-box; }}
body {{ font-family: Calibri, Carlito, Arial, sans-serif; font-size: {BODY_PT}pt; color: #111;
       line-height: 1.2; margin: 0; }}
h1 {{ font-size: 19pt; margin: 0; line-height: 1.1; }}
.headline {{ font-size: 11pt; color: #444; margin: 0; }}
.contact {{ font-size: 9pt; color: #444; margin: 0; }}
h2 {{ font-size: 10.2pt; text-transform: uppercase; letter-spacing: .04em; margin: 6pt 0 2pt;
     padding-bottom: 1pt; border-bottom: 0.75pt solid #999; }}
p {{ margin: 0; }}
.row {{ display: flex; justify-content: space-between; font-weight: 700; margin-top: 4pt; }}
.org {{ font-style: italic; color: #444; margin-bottom: 1pt; }}
ul {{ margin: 0; padding-left: 13pt; }}
li {{ margin-bottom: 1pt; }}
.skill {{ margin-bottom: 1.5pt; }}
.proj {{ margin-bottom: 1.5pt; }}
.muted {{ color: #444; }}
a {{ color: #1f4e8c; text-decoration: none; }}
.edu {{ display: flex; justify-content: space-between; }}
.cert {{ margin-top: 1pt; }}
"""


def build_html():
    e = html.escape
    out = [f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>{e(NAME)} Resume</title>"
           f"<style>{CSS}</style></head><body>"]
    out.append(f"<h1>{e(NAME)}</h1><p class='headline'>{e(HEADLINE)}</p>")
    out += ["<p class='contact'>" + " | ".join(f"<a href='{e(u)}'>{e(t)}</a>" for t, u in line) + "</p>"
            for line in CONTACT]
    out.append(f"<h2>Summary</h2><p>{e(SUMMARY)}</p>")
    out.append("<h2>Skills</h2>")
    out += [f"<p class='skill'><b>{e(g)}:</b> {e(items)}</p>" for g, items in SKILLS]
    out.append("<h2>Experience</h2>")
    for title, dates, org, bullets in JOBS:
        out.append(f"<div class='row'><span>{e(title)}</span><span>{e(dates)}</span></div>"
                   f"<p class='org'>{e(org)}</p><ul>")
        out += [f"<li>{e(b)}</li>" for b in bullets]
        out.append("</ul>")
    out.append("<h2>Projects</h2>")
    out += [f"<p class='proj'><b>{e(n)}:</b> {e(d)} <a href='{e(url)}'>{e(text)}</a></p>"
            for n, d, (text, url) in PROJECTS]
    out.append("<h2>Education &amp; certifications</h2>" + "".join(
        f"<div class='edu'><span>{e(d)}</span><span>{e(y)}</span></div>" for d, y in EDUCATION))
    out.append(f"<p class='cert'>{e(CERTS)}</p>")
    out.append("</body></html>")
    return "".join(out)


def build_pdf(html_path, pdf_path):
    html_path.write_text(build_html(), encoding="utf-8")
    subprocess.run(
        [str(EDGE), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
         f"--print-to-pdf={pdf_path}", html_path.as_uri()],
        check=True, capture_output=True, timeout=120,
    )


# ---------------------------------------------------------------- plain text (paste-in forms)

def build_text():
    lines = [NAME, HEADLINE, " | ".join(t for line in CONTACT for t, _ in line), "",
             "SUMMARY", SUMMARY, "", "SKILLS"]
    lines += [f"{g}: {items}" for g, items in SKILLS]
    lines += ["", "EXPERIENCE"]
    for title, dates, org, bullets in JOBS:
        lines += ["", f"{title} | {dates}", org] + [f"- {b}" for b in bullets]
    lines += ["", "PROJECTS"] + [f"- {n}: {d} {url}" for n, d, (_, url) in PROJECTS]
    lines += ["", "EDUCATION & CERTIFICATIONS"] + [f"{d}, {y}" for d, y in EDUCATION] + [CERTS]
    return "\n".join(lines) + "\n"


def page_count(pdf_path):
    return subprocess.run(["pdftotext", str(pdf_path), "-"], capture_output=True,
                          text=True, encoding="utf-8", errors="replace").stdout.count("\f")


def main():
    build = HERE / "build"
    build.mkdir(exist_ok=True)
    stem = "Siddhesh_Amrale_Resume"
    pdf = HERE / f"{stem}.pdf"
    build_docx(HERE / f"{stem}.docx")
    (HERE / f"{stem}.txt").write_text(build_text(), encoding="utf-8")
    build_pdf(build / f"{stem}.html", pdf)
    pages = page_count(pdf)
    if pages != 1:
        print(f"resume is {pages} pages; it must fit on one", file=sys.stderr)
        return 1
    # The site's Resume button serves this same file.
    shutil.copyfile(pdf, HERE.parent / "public" / f"{stem}.pdf")
    print("built", stem, "(1 page)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
